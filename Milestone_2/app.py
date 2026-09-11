import streamlit as st
import os
import ctypes
import sys
import tempfile
import shutil
import whisper
from moviepy import VideoFileClip
from google import genai
import chromadb
from sentence_transformers import SentenceTransformer
import soundfile as sf
import torch
from pyannote.audio import Pipeline
import io
import markdown
import docx
from xhtml2pdf import pisa
import re
import time
from database import init_db, create_user, authenticate_user, get_user_meetings, get_meeting_by_id

# Initialize database on startup
init_db()

# --- WORKAROUNDS ---
try:
    pytorch_lib_dir = os.path.join(sys.prefix, "Lib", "site-packages", "torch", "lib")
    c10_dll = os.path.join(pytorch_lib_dir, "c10.dll")
    if os.path.exists(c10_dll):
        ctypes.CDLL(c10_dll)
except Exception:
    pass

try:
    import imageio_ffmpeg
    ffmpeg_src = imageio_ffmpeg.get_ffmpeg_exe()
    ffmpeg_dst = os.path.join(os.getcwd(), "ffmpeg.exe")
    if not os.path.exists(ffmpeg_dst):
        shutil.copy(ffmpeg_src, ffmpeg_dst)
except Exception:
    pass
# -------------------

st.set_page_config(
    page_title="SynthAI — Meeting Intelligence",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ─────────────────────────────────────────────
#  DESIGN SYSTEM — injected CSS
# ─────────────────────────────────────────────
st.markdown("""
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500&display=swap" rel="stylesheet">
<link href="https://fonts.googleapis.com/css2?family=Material+Symbols+Rounded:opsz,wght,FILL,GRAD@20..48,100..700,0..1,-50..200" rel="stylesheet">

<style>
/* ── ROOT VARIABLES ─────────────────────────────── */
:root {
  --bg-void:       #060812;
  --bg-base:       #0A0E1A;
  --bg-surface:    #0F1628;
  --bg-elevated:   #151c35;
  --glass-bg:      rgba(15, 22, 40, 0.65);
  --glass-border:  rgba(99, 102, 241, 0.18);
  --glass-border-hover: rgba(99, 102, 241, 0.45);

  --indigo:        #6366F1;
  --indigo-light:  #818CF8;
  --indigo-dark:   #4F46E5;
  --cyan:          #22D3EE;
  --cyan-light:    #67E8F9;
  --emerald:       #10B981;
  --amber:         #F59E0B;
  --rose:          #F43F5E;

  --text-primary:   #F0F4FF;
  --text-secondary: #94A3B8;
  --text-muted:     #4E5E7A;

  --radius-sm:  6px;
  --radius-md:  12px;
  --radius-lg:  20px;
  --radius-xl:  28px;

  --shadow-glow-indigo: 0 0 30px rgba(99, 102, 241, 0.25);
  --shadow-glow-cyan:   0 0 30px rgba(34, 211, 238, 0.20);
  --shadow-card:        0 4px 32px rgba(0,0,0,0.45);
  --transition: all 0.25s cubic-bezier(0.4, 0, 0.2, 1);
}

/* ── GLOBAL RESET ───────────────────────────────── */
*, *::before, *::after { box-sizing: border-box; }

html, body, [data-testid="stAppViewContainer"],
[data-testid="stApp"] {
  background: var(--bg-void) !important;
  font-family: 'Inter', sans-serif !important;
  color: var(--text-primary) !important;
}

[data-testid="stAppViewContainer"] {
  background:
    radial-gradient(ellipse 80% 60% at 20% -10%, rgba(99,102,241,0.12) 0%, transparent 60%),
    radial-gradient(ellipse 60% 40% at 80% 110%, rgba(34,211,238,0.08) 0%, transparent 55%),
    var(--bg-void) !important;
  min-height: 100vh;
}

[data-testid="stMain"] {
  background: transparent !important;
}

/* ── SIDEBAR ────────────────────────────────────── */
[data-testid="stSidebar"] {
  background: rgba(10, 14, 26, 0.92) !important;
  border-right: 1px solid var(--glass-border) !important;
  backdrop-filter: blur(20px) !important;
}

[data-testid="stSidebar"] > div {
  padding-top: 1.5rem !important;
}

/* sidebar title */
[data-testid="stSidebar"] h1 {
  font-size: 1rem !important;
  font-weight: 700 !important;
  letter-spacing: 0.05em !important;
  text-transform: uppercase !important;
  color: var(--indigo-light) !important;
  padding: 0 1rem 0.5rem !important;
  border-bottom: 1px solid var(--glass-border) !important;
  margin-bottom: 1.25rem !important;
}

/* sidebar labels */
[data-testid="stSidebar"] label {
  font-size: 0.72rem !important;
  font-weight: 600 !important;
  letter-spacing: 0.08em !important;
  text-transform: uppercase !important;
  color: var(--text-secondary) !important;
}

/* ── INPUTS & TEXT FIELDS (Global, Sidebar & Forms) ───────────────── */
[data-testid="stTextInput"] [data-baseweb="base-input"],
[data-testid="stTextInput"] [data-baseweb="input"],
[data-testid="stTextInput"] div[data-baseweb="base-input"] > div,
[data-testid="stTextInput"] div[data-baseweb="input"] > div,
[data-testid="stForm"] [data-baseweb="base-input"],
[data-testid="stForm"] [data-baseweb="input"],
[data-testid="stForm"] div[data-baseweb="base-input"] > div,
[data-testid="stForm"] div[data-baseweb="input"] > div,
div[data-baseweb="base-input"],
div[data-baseweb="input"] {
  background-color: rgba(15, 22, 40, 0.85) !important;
  background: rgba(15, 22, 40, 0.85) !important;
  border: 1px solid var(--glass-border) !important;
  border-radius: var(--radius-sm) !important;
  color: var(--text-primary) !important;
}

[data-testid="stSidebar"] [data-baseweb="base-input"],
[data-testid="stSidebar"] [data-baseweb="input"],
[data-testid="stSidebar"] div[data-baseweb="base-input"] > div,
[data-testid="stSidebar"] div[data-baseweb="input"] > div {
  background-color: rgba(99, 102, 241, 0.08) !important;
  background: rgba(99, 102, 241, 0.08) !important;
  border: 1px solid var(--glass-border) !important;
  border-radius: var(--radius-sm) !important;
}

[data-testid="stTextInput"] input,
[data-testid="stForm"] input,
[data-testid="stSidebar"] input,
input[type="text"],
input[type="password"] {
  background: transparent !important;
  background-color: transparent !important;
  color: #F0F4FF !important;
  font-family: 'Inter', sans-serif !important;
  font-size: 0.88rem !important;
  caret-color: #F0F4FF !important;
  border: none !important;
}

[data-testid="stSidebar"] input {
  font-family: 'JetBrains Mono', monospace !important;
  font-size: 0.8rem !important;
}

/* Visibility toggle button in password inputs */
[data-baseweb="base-input"] button,
[data-testid="stTextInput"] button,
[data-testid="stSidebar"] button[kind="secondary"] {
  background: transparent !important;
  background-color: transparent !important;
  color: var(--text-secondary) !important;
  border: none !important;
  box-shadow: none !important;
}

input:focus, [data-baseweb="base-input"]:focus-within {
  border-color: var(--indigo) !important;
  box-shadow: 0 0 0 2px rgba(99, 102, 241, 0.25) !important;
}

/* ── HEADINGS & TEXT ────────────────────────────── */
h1, h2, h3, h4 {
  font-family: 'Inter', sans-serif !important;
  color: var(--text-primary) !important;
  letter-spacing: -0.02em !important;
}

p, li, div {
  font-family: 'Inter', sans-serif !important;
}

span:not([class*="material"]):not([data-testid*="Icon"]) {
  font-family: 'Inter', sans-serif !important;
}

/* Material Symbols & Icons */
.material-symbols-rounded,
.material-symbols-outlined,
.material-icons,
[data-testid="stIconMaterial"],
[data-baseweb="base-input"] button span,
[data-testid="stTextInput"] button span {
  font-family: 'Material Symbols Rounded', 'Material Icons', sans-serif !important;
  font-size: 1.2rem !important;
  line-height: 1 !important;
  letter-spacing: normal !important;
  text-transform: none !important;
  display: inline-block !important;
  white-space: nowrap !important;
  word-wrap: normal !important;
  direction: ltr !important;
  -webkit-font-smoothing: antialiased !important;
}

/* ── HERO BANNER ────────────────────────────────── */
.hero-container {
  text-align: center;
  padding: 3.5rem 2rem 2rem;
  animation: fadeInDown 0.7s ease-out both;
}

.hero-logo {
  display: inline-flex;
  align-items: center;
  gap: 0.65rem;
  margin-bottom: 1rem;
}

.hero-logo-icon {
  width: 52px; height: 52px;
  background: linear-gradient(135deg, var(--indigo) 0%, var(--cyan) 100%);
  border-radius: 14px;
  display: flex; align-items: center; justify-content: center;
  font-size: 1.6rem;
  box-shadow: var(--shadow-glow-indigo);
}

.hero-brand {
  font-size: 2.4rem;
  font-weight: 800;
  background: linear-gradient(135deg, #F0F4FF 0%, var(--indigo-light) 45%, var(--cyan) 100%);
  -webkit-background-clip: text;
  -webkit-text-fill-color: transparent;
  background-clip: text;
  letter-spacing: -0.03em;
  line-height: 1;
}

.hero-tagline {
  font-size: 1.05rem;
  color: var(--text-secondary);
  font-weight: 400;
  margin-top: 0.5rem;
  letter-spacing: 0.01em;
}

.hero-badges {
  display: flex;
  gap: 0.5rem;
  justify-content: center;
  flex-wrap: wrap;
  margin-top: 1.25rem;
}

.hero-badge {
  display: inline-flex;
  align-items: center;
  gap: 0.35rem;
  padding: 0.3rem 0.75rem;
  border-radius: 100px;
  font-size: 0.72rem;
  font-weight: 600;
  letter-spacing: 0.04em;
  text-transform: uppercase;
  border: 1px solid;
}

.badge-whisper { background: rgba(99,102,241,0.1); border-color: rgba(99,102,241,0.3); color: var(--indigo-light); }
.badge-gemini  { background: rgba(34,211,238,0.1); border-color: rgba(34,211,238,0.3); color: var(--cyan-light); }
.badge-rag     { background: rgba(16,185,129,0.1); border-color: rgba(16,185,129,0.3); color: #6EE7B7; }

/* ── GLASS CARD ─────────────────────────────────── */
.glass-card {
  background: var(--glass-bg);
  border: 1px solid var(--glass-border);
  border-radius: var(--radius-lg);
  backdrop-filter: blur(20px);
  -webkit-backdrop-filter: blur(20px);
  padding: 1.75rem;
  box-shadow: var(--shadow-card);
  transition: var(--transition);
  animation: fadeInUp 0.5s ease-out both;
}

.glass-card:hover {
  border-color: var(--glass-border-hover);
  box-shadow: var(--shadow-glow-indigo), var(--shadow-card);
  transform: translateY(-2px);
}

/* ── UPLOAD ZONE ────────────────────────────────── */
.upload-zone-wrapper {
  background: linear-gradient(135deg, rgba(99,102,241,0.06) 0%, rgba(34,211,238,0.04) 100%);
  border: 2px dashed rgba(99,102,241,0.35);
  border-radius: var(--radius-xl);
  padding: 2.5rem 2rem;
  text-align: center;
  transition: var(--transition);
  position: relative;
  overflow: hidden;
  margin-bottom: 1rem;
}

.upload-zone-wrapper::before {
  content: '';
  position: absolute;
  inset: 0;
  background: linear-gradient(90deg, transparent, rgba(99,102,241,0.04), transparent);
  animation: shimmer 3s ease-in-out infinite;
}

.upload-icon {
  font-size: 3rem;
  margin-bottom: 0.75rem;
  display: block;
  animation: float 3s ease-in-out infinite;
}

.upload-title {
  font-size: 1.15rem;
  font-weight: 600;
  color: var(--text-primary);
  margin-bottom: 0.35rem;
}

.upload-subtitle {
  font-size: 0.82rem;
  color: var(--text-muted);
}

.format-chips {
  display: flex;
  gap: 0.4rem;
  justify-content: center;
  margin-top: 0.75rem;
  flex-wrap: wrap;
}

.format-chip {
  padding: 0.2rem 0.6rem;
  background: rgba(99,102,241,0.12);
  border: 1px solid rgba(99,102,241,0.25);
  border-radius: 100px;
  font-size: 0.7rem;
  font-weight: 600;
  color: var(--indigo-light);
  letter-spacing: 0.05em;
  font-family: 'JetBrains Mono', monospace;
}

/* ── STREAMLIT FILE UPLOADER OVERRIDE ───────────── */
[data-testid="stFileUploader"] {
  background: transparent !important;
}

[data-testid="stFileUploader"] > div {
  background: rgba(99, 102, 241, 0.05) !important;
  border: 2px dashed rgba(99, 102, 241, 0.35) !important;
  border-radius: var(--radius-lg) !important;
  transition: var(--transition) !important;
}

[data-testid="stFileUploader"] > div:hover {
  border-color: var(--indigo) !important;
  background: rgba(99, 102, 241, 0.10) !important;
}

[data-testid="stFileUploader"] label {
  color: var(--text-secondary) !important;
  font-size: 0.85rem !important;
}

[data-testid="stFileDropzone"] {
  background: transparent !important;
}

/* ── BUTTONS ────────────────────────────────────── */
[data-testid="stButton"] > button,
.stButton > button,
[data-testid="stFormSubmitButton"] > button,
button[kind="secondaryFormSubmit"],
button[kind="primaryFormSubmit"] {
  background: linear-gradient(135deg, var(--indigo-dark) 0%, var(--indigo) 50%, #7C3AED 100%) !important;
  color: #fff !important;
  border: none !important;
  border-radius: var(--radius-md) !important;
  font-family: 'Inter', sans-serif !important;
  font-weight: 600 !important;
  font-size: 0.88rem !important;
  padding: 0.65rem 1.5rem !important;
  letter-spacing: 0.02em !important;
  cursor: pointer !important;
  transition: var(--transition) !important;
  box-shadow: 0 4px 20px rgba(99, 102, 241, 0.35) !important;
  position: relative !important;
  overflow: hidden !important;
}

[data-testid="stButton"] > button::after,
[data-testid="stFormSubmitButton"] > button::after {
  content: '';
  position: absolute;
  inset: 0;
  background: linear-gradient(135deg, rgba(255,255,255,0.1), transparent);
  opacity: 0;
  transition: var(--transition);
}

[data-testid="stButton"] > button:hover,
[data-testid="stFormSubmitButton"] > button:hover {
  transform: translateY(-2px) !important;
  box-shadow: 0 8px 30px rgba(99, 102, 241, 0.55) !important;
}

[data-testid="stButton"] > button:hover::after,
[data-testid="stFormSubmitButton"] > button:hover::after {
  opacity: 1 !important;
}

[data-testid="stButton"] > button:active,
[data-testid="stFormSubmitButton"] > button:active {
  transform: translateY(0) !important;
}

/* ── DOWNLOAD BUTTONS ───────────────────────────── */
[data-testid="stDownloadButton"] > button {
  background: rgba(15, 22, 40, 0.8) !important;
  color: var(--text-primary) !important;
  border: 1px solid var(--glass-border) !important;
  border-radius: var(--radius-md) !important;
  font-family: 'Inter', sans-serif !important;
  font-weight: 500 !important;
  font-size: 0.85rem !important;
  padding: 0.6rem 1.25rem !important;
  transition: var(--transition) !important;
  width: 100% !important;
}

[data-testid="stDownloadButton"] > button:hover {
  border-color: var(--indigo) !important;
  background: rgba(99, 102, 241, 0.12) !important;
  transform: translateY(-1px) !important;
  box-shadow: var(--shadow-glow-indigo) !important;
}

/* ── TABS ───────────────────────────────────────── */
[data-testid="stTabs"] {
  background: transparent !important;
}

[data-testid="stTabsList"],
div[data-baseweb="tab-list"],
div[role="tablist"] {
  background: rgba(10, 14, 26, 0.7) !important;
  border: 1px solid var(--glass-border) !important;
  border-radius: var(--radius-md) !important;
  padding: 4px !important;
  gap: 4px !important;
  backdrop-filter: blur(12px) !important;
  width: fit-content !important;
}

div[data-baseweb="tab-highlight"],
div[data-baseweb="tab-border"] {
  display: none !important;
}

button[data-baseweb="tab"],
button[role="tab"],
[data-testid="stTab"] {
  background: transparent !important;
  color: var(--text-muted) !important;
  border-radius: 8px !important;
  font-family: 'Inter', sans-serif !important;
  font-size: 0.85rem !important;
  font-weight: 500 !important;
  padding: 0.5rem 1.25rem !important;
  border: none !important;
  transition: var(--transition) !important;
  letter-spacing: 0.01em !important;
}

button[data-baseweb="tab"][aria-selected="true"],
button[role="tab"][aria-selected="true"],
[data-testid="stTab"][aria-selected="true"] {
  background: linear-gradient(135deg, var(--indigo-dark), var(--indigo)) !important;
  color: #fff !important;
  box-shadow: 0 2px 12px rgba(99,102,241,0.4) !important;
}

button[data-baseweb="tab"]:hover:not([aria-selected="true"]),
button[role="tab"]:hover:not([aria-selected="true"]),
[data-testid="stTab"]:hover:not([aria-selected="true"]) {
  color: var(--text-secondary) !important;
  background: rgba(99,102,241,0.08) !important;
}

[data-testid="stTabsContent"] {
  background: transparent !important;
  border: none !important;
  padding: 1.5rem 0 0 !important;
}

/* ── PROGRESS / STEPS ───────────────────────────── */
.steps-container {
  display: flex;
  justify-content: center;
  align-items: flex-start;
  gap: 0;
  margin: 1.5rem 0;
  padding: 1.5rem;
  background: var(--glass-bg);
  border: 1px solid var(--glass-border);
  border-radius: var(--radius-lg);
  backdrop-filter: blur(20px);
}

.step-item {
  display: flex;
  flex-direction: column;
  align-items: center;
  flex: 1;
  position: relative;
}

.step-item:not(:last-child)::after {
  content: '';
  position: absolute;
  top: 20px;
  left: 50%;
  width: 100%;
  height: 2px;
  background: var(--glass-border);
  z-index: 0;
}

.step-item.active:not(:last-child)::after {
  background: linear-gradient(90deg, var(--indigo), transparent);
}

.step-item.done:not(:last-child)::after {
  background: var(--indigo);
}

.step-dot {
  width: 40px; height: 40px;
  border-radius: 50%;
  display: flex; align-items: center; justify-content: center;
  font-size: 1rem;
  position: relative; z-index: 1;
  transition: var(--transition);
}

.step-dot.waiting {
  background: rgba(78, 94, 122, 0.3);
  border: 2px solid var(--text-muted);
}

.step-dot.active {
  background: linear-gradient(135deg, var(--indigo-dark), var(--indigo));
  border: 2px solid var(--indigo-light);
  box-shadow: var(--shadow-glow-indigo);
  animation: pulse-glow 1.5s ease-in-out infinite;
}

.step-dot.done {
  background: linear-gradient(135deg, var(--emerald), #059669);
  border: 2px solid #6EE7B7;
}

.step-label {
  margin-top: 0.5rem;
  font-size: 0.7rem;
  font-weight: 600;
  letter-spacing: 0.05em;
  text-transform: uppercase;
  text-align: center;
  color: var(--text-muted);
}

.step-label.active { color: var(--indigo-light); }
.step-label.done   { color: var(--emerald); }

/* ── SECTION HEADERS ────────────────────────────── */
.section-header {
  display: flex;
  align-items: center;
  gap: 0.65rem;
  margin-bottom: 1.25rem;
}

.section-header-icon {
  width: 36px; height: 36px;
  border-radius: 10px;
  display: flex; align-items: center; justify-content: center;
  font-size: 1rem;
  flex-shrink: 0;
}

.icon-indigo { background: rgba(99,102,241,0.15); }
.icon-cyan   { background: rgba(34,211,238,0.12); }
.icon-emerald{ background: rgba(16,185,129,0.12); }
.icon-amber  { background: rgba(245,158,11,0.12); }

.section-header-text h3 {
  font-size: 1rem !important;
  font-weight: 700 !important;
  color: var(--text-primary) !important;
  margin: 0 !important;
}

.section-header-text span {
  font-size: 0.75rem;
  color: var(--text-muted);
  font-weight: 400;
}

/* ── REPORT CARDS ───────────────────────────────── */
.report-executive {
  background: linear-gradient(135deg, rgba(99,102,241,0.1) 0%, rgba(34,211,238,0.06) 100%);
  border: 1px solid rgba(99,102,241,0.25);
  border-radius: var(--radius-lg);
  padding: 1.75rem;
  margin-bottom: 1.25rem;
  position: relative;
  overflow: hidden;
}

.report-executive::before {
  content: '';
  position: absolute;
  top: 0; left: 0;
  width: 4px; height: 100%;
  background: linear-gradient(180deg, var(--indigo), var(--cyan));
}

.report-executive p, .report-executive ul, .report-executive li {
  color: var(--text-secondary) !important;
  line-height: 1.7 !important;
}

/* ── CHAT ───────────────────────────────────────── */
[data-testid="stChatMessage"] {
  background: transparent !important;
  border: none !important;
  padding: 0.5rem 0 !important;
}

[data-testid="stChatMessage"][data-testid*="user"] .stMarkdown,
[data-testid="stChatMessageContent"] {
  background: transparent !important;
}

.chat-message-user {
  background: linear-gradient(135deg, rgba(99,102,241,0.2), rgba(99,102,241,0.12)) !important;
  border: 1px solid rgba(99,102,241,0.25) !important;
  border-radius: 18px 18px 4px 18px !important;
  padding: 0.85rem 1.1rem !important;
  margin-left: auto !important;
  max-width: 80% !important;
  color: var(--text-primary) !important;
}

.chat-message-ai {
  background: rgba(15,22,40,0.7) !important;
  border: 1px solid var(--glass-border) !important;
  border-radius: 18px 18px 18px 4px !important;
  padding: 0.85rem 1.1rem !important;
  max-width: 90% !important;
  color: var(--text-primary) !important;
  backdrop-filter: blur(8px) !important;
}

/* Streamlit chat input */
[data-testid="stChatInput"] {
  background: rgba(15, 22, 40, 0.9) !important;
  background-color: rgba(15, 22, 40, 0.9) !important;
  border: 1px solid var(--glass-border) !important;
  border-radius: var(--radius-lg) !important;
  backdrop-filter: blur(16px) !important;
}

[data-testid="stChatInput"] > div,
[data-testid="stChatInput"] form,
[data-testid="stChatInput"] [data-baseweb="base-input"],
[data-testid="stChatInput"] [data-baseweb="textarea"],
[data-testid="stChatInput"] [data-baseweb="base-input"] > div,
[data-testid="stChatInput"] [data-baseweb="textarea"] > div {
  background: transparent !important;
  background-color: transparent !important;
  border-color: transparent !important;
}

[data-testid="stChatInput"] textarea {
  background: transparent !important;
  background-color: transparent !important;
  color: #F0F4FF !important;
  font-family: 'Inter', sans-serif !important;
  font-size: 0.9rem !important;
  caret-color: #F0F4FF !important;
}

[data-testid="stChatInput"] textarea::placeholder {
  color: #64748B !important;
}

[data-testid="stChatInput"] button {
  background: linear-gradient(135deg, var(--indigo-dark) 0%, var(--indigo) 100%) !important;
  color: #ffffff !important;
  border: none !important;
  border-radius: 8px !important;
}

/* Ensure the bottom-docked chat container is also styled */
[data-testid="stBottom"],
[data-testid="stBottom"] > div {
  background: var(--bg-void) !important;
  background-color: var(--bg-void) !important;
}

/* ── TRANSCRIPT CODE BLOCK ──────────────────────── */
[data-testid="stCode"],
.stCode {
  background: rgba(6, 8, 18, 0.8) !important;
  border: 1px solid var(--glass-border) !important;
  border-radius: var(--radius-md) !important;
}

pre, code {
  font-family: 'JetBrains Mono', monospace !important;
  font-size: 0.78rem !important;
  color: #CBD5E1 !important;
  background: transparent !important;
}

/* ── ALERTS / INFO BANNERS ──────────────────────── */
[data-testid="stAlert"] {
  border-radius: var(--radius-md) !important;
  border: 1px solid !important;
  backdrop-filter: blur(12px) !important;
}

[data-testid="stAlert"][data-baseweb="notification"] {
  background: rgba(15, 22, 40, 0.8) !important;
}

.stSuccess {
  background: rgba(16, 185, 129, 0.1) !important;
  border-color: rgba(16, 185, 129, 0.3) !important;
}

.stInfo {
  background: rgba(99, 102, 241, 0.1) !important;
  border-color: rgba(99, 102, 241, 0.3) !important;
}

.stWarning {
  background: rgba(245, 158, 11, 0.1) !important;
  border-color: rgba(245, 158, 11, 0.3) !important;
}

.stError {
  background: rgba(244, 63, 94, 0.1) !important;
  border-color: rgba(244, 63, 94, 0.3) !important;
}

/* ── SPINNER OVERRIDE ───────────────────────────── */
[data-testid="stSpinner"] > div {
  border-color: var(--indigo) transparent transparent transparent !important;
}

/* ── DIVIDERS ───────────────────────────────────── */
hr {
  border: none !important;
  border-top: 1px solid var(--glass-border) !important;
  margin: 1.5rem 0 !important;
}

/* ── SCROLLBAR ──────────────────────────────────── */
::-webkit-scrollbar { width: 6px; height: 6px; }
::-webkit-scrollbar-track { background: transparent; }
::-webkit-scrollbar-thumb {
  background: rgba(99,102,241,0.3);
  border-radius: 100px;
}
::-webkit-scrollbar-thumb:hover { background: var(--indigo); }

/* ── METRIC CARDS ───────────────────────────────── */
.metric-card {
  background: var(--glass-bg);
  border: 1px solid var(--glass-border);
  border-radius: var(--radius-md);
  padding: 1.1rem 1.25rem;
  display: flex;
  align-items: center;
  gap: 0.85rem;
  transition: var(--transition);
}

.metric-card:hover {
  border-color: var(--glass-border-hover);
  transform: translateY(-1px);
}

.metric-icon {
  font-size: 1.5rem;
  width: 44px; height: 44px;
  display: flex; align-items: center; justify-content: center;
  border-radius: 10px;
  flex-shrink: 0;
}

.metric-label {
  font-size: 0.68rem;
  font-weight: 600;
  letter-spacing: 0.07em;
  text-transform: uppercase;
  color: var(--text-muted);
  margin-bottom: 0.15rem;
}

.metric-value {
  font-size: 1.15rem;
  font-weight: 700;
  color: var(--text-primary);
  letter-spacing: -0.01em;
}

/* ── STATUS DOT ─────────────────────────────────── */
.status-connected {
  display: inline-flex;
  align-items: center;
  gap: 0.4rem;
  padding: 0.3rem 0.7rem;
  background: rgba(16, 185, 129, 0.1);
  border: 1px solid rgba(16, 185, 129, 0.3);
  border-radius: 100px;
  font-size: 0.7rem;
  font-weight: 600;
  color: #6EE7B7;
  margin-top: 0.5rem;
}

.status-dot {
  width: 6px; height: 6px;
  border-radius: 50%;
  background: var(--emerald);
  animation: pulse-dot 2s ease-in-out infinite;
}

/* ── EXPORT SECTION ─────────────────────────────── */
.export-header {
  font-size: 0.72rem;
  font-weight: 700;
  letter-spacing: 0.09em;
  text-transform: uppercase;
  color: var(--text-muted);
  margin-bottom: 0.75rem;
  display: flex;
  align-items: center;
  gap: 0.4rem;
}

/* ── ANIMATIONS ─────────────────────────────────── */
@keyframes fadeInDown {
  from { opacity: 0; transform: translateY(-20px); }
  to   { opacity: 1; transform: translateY(0); }
}

@keyframes fadeInUp {
  from { opacity: 0; transform: translateY(20px); }
  to   { opacity: 1; transform: translateY(0); }
}

@keyframes float {
  0%, 100% { transform: translateY(0); }
  50%       { transform: translateY(-6px); }
}

@keyframes pulse-glow {
  0%, 100% { box-shadow: 0 0 0 0 rgba(99,102,241,0.5); }
  50%       { box-shadow: 0 0 0 8px rgba(99,102,241,0); }
}

@keyframes pulse-dot {
  0%, 100% { opacity: 1; }
  50%       { opacity: 0.4; }
}

@keyframes shimmer {
  0%   { transform: translateX(-100%); }
  50%  { transform: translateX(100%); }
  100% { transform: translateX(100%); }
}

@keyframes spin-glow {
  0%   { transform: rotate(0deg); }
  100% { transform: rotate(360deg); }
}

/* ── HIDE STREAMLIT CHROME ──────────────────────── */
#MainMenu, footer { visibility: hidden !important; }
header { background: transparent !important; }
[data-testid="stDecoration"] { display: none !important; }
[data-testid="stToolbar"]  { display: none !important; }

/* ── FORCE SIDEBAR ALWAYS VISIBLE ──────────────── */
[data-testid="stSidebar"] {
  min-width: 280px !important;
  max-width: 320px !important;
  width: 300px !important;
  transform: none !important;
  visibility: visible !important;
  display: block !important;
}
[data-testid="collapsedControl"] {
  display: block !important;
  visibility: visible !important;
  color: white !important;
}

/* ── MARKDOWN INSIDE GLASS CARDS ────────────────── */
.element-container .stMarkdown p { color: var(--text-secondary) !important; line-height: 1.75 !important; }
.element-container .stMarkdown h2 { color: var(--text-primary) !important; margin-top: 1rem !important; }
.element-container .stMarkdown h3 { color: var(--indigo-light) !important; }
.element-container .stMarkdown li { color: var(--text-secondary) !important; }
.element-container .stMarkdown strong { color: var(--text-primary) !important; }

/* ── COLUMN GAPS ────────────────────────────────── */
[data-testid="stHorizontalBlock"] { gap: 1rem !important; }

/* ── PROCESS NEW FILE btn ───────────────────────── */
.new-file-btn > button {
  background: rgba(244, 63, 94, 0.08) !important;
  border: 1px solid rgba(244, 63, 94, 0.3) !important;
  color: #FDA4AF !important;
  box-shadow: none !important;
}

.new-file-btn > button:hover {
  background: rgba(244, 63, 94, 0.15) !important;
  border-color: rgba(244, 63, 94, 0.6) !important;
  transform: translateY(-1px) !important;
  box-shadow: 0 0 20px rgba(244, 63, 94, 0.2) !important;
}

/* selectbox */
[data-testid="stSelectbox"] select,
[data-baseweb="select"] {
  background: rgba(15, 22, 40, 0.8) !important;
  border-color: var(--glass-border) !important;
  color: var(--text-primary) !important;
}
</style>
""", unsafe_allow_html=True)


# ─────────────────────────────────────────────
#  SESSION STATE
# ─────────────────────────────────────────────
for key, default in [
    ("transcript", None),
    ("chroma_collection", None),
    ("messages", []),
    ("report_summary", None),
    ("structured_data", None),
    ("user_id", None),
    ("username", None),
    ("current_page", "Command Center"),
    ("active_meeting_id", None),
    ("current_meeting_data", None),
]:
    if key not in st.session_state:
        st.session_state[key] = default


# ─────────────────────────────────────────────
#  AUTHENTICATION (before sidebar)
# ─────────────────────────────────────────────
if not st.session_state.user_id:
    # Hide sidebar when not logged in
    st.markdown('<style>[data-testid="stSidebar"] { display: none !important; } [data-testid="collapsedControl"] { display: none !important; }</style>', unsafe_allow_html=True)

    st.markdown("""
    <div class="hero-container">
      <div class="hero-logo">
        <div class="hero-logo-icon">🧠</div>
        <div class="hero-brand">SynthAI</div>
      </div>
      <div class="hero-tagline">Please log in to access your meeting intelligence.</div>
    </div>
    """, unsafe_allow_html=True)
    
    auth_col1, auth_col2, auth_col3 = st.columns([1, 2, 1])
    with auth_col2:
        tab_login, tab_signup = st.tabs(["Login", "Sign Up"])
        
        with tab_login:
            with st.form("login_form"):
                log_user = st.text_input("Username")
                log_pass = st.text_input("Password", type="password")
                if st.form_submit_button("Login", use_container_width=True):
                    from database import authenticate_user
                    user = authenticate_user(log_user, log_pass)
                    if user:
                        st.session_state.user_id = user["id"]
                        st.session_state.username = user["username"]
                        st.rerun()
                    else:
                        st.error("Invalid username or password.")
                        
        with tab_signup:
            with st.form("signup_form"):
                reg_user = st.text_input("New Username")
                reg_pass = st.text_input("New Password", type="password")
                if st.form_submit_button("Sign Up", use_container_width=True):
                    from database import create_user
                    if len(reg_user.strip()) < 3:
                        st.error("Username must be at least 3 characters.")
                    elif len(reg_pass) < 6:
                        st.error("Password must be at least 6 characters.")
                    else:
                        success = create_user(reg_user, reg_pass)
                        if success:
                            st.success("Account created! You can now log in.")
                        else:
                            st.error("Username already exists.")
    st.stop()


# ─────────────────────────────────────────────
#  SIDEBAR
# ─────────────────────────────────────────────
with st.sidebar:
    st.markdown("""
    <div style="display:flex;align-items:center;gap:0.5rem;padding:0 1rem 1rem;">
      <div style="width:32px;height:32px;background:linear-gradient(135deg,#6366F1,#22D3EE);
                  border-radius:9px;display:flex;align-items:center;justify-content:center;
                  font-size:1.1rem;">🧠</div>
      <div style="font-size:1rem;font-weight:800;background:linear-gradient(135deg,#F0F4FF,#818CF8,#22D3EE);
                  -webkit-background-clip:text;-webkit-text-fill-color:transparent;
                  background-clip:text;letter-spacing:-0.02em;">SynthAI</div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown('<p style="font-size:0.68rem;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;color:#4E5E7A;padding:0 0 0.5rem;">API Configuration</p>', unsafe_allow_html=True)

    gemini_key = st.text_input(
        "Gemini API Key",
        type="password",
        value=os.getenv("GEMINI_API_KEY", ""),
        help="Your Google Gemini API key for report generation and chat",
    )
    hf_token = st.text_input(
        "Hugging Face Token",
        type="password",
        value=os.getenv("HF_TOKEN", ""),
        help="Required for Pyannote speaker diarization model",
    )

    keys_ok = bool(gemini_key and hf_token)

    if keys_ok:
        st.markdown("""
        <div class="status-connected">
          <div class="status-dot"></div>Keys Configured
        </div>
        """, unsafe_allow_html=True)
    else:
        st.markdown("""
        <div style="display:inline-flex;align-items:center;gap:0.4rem;padding:0.3rem 0.7rem;
                    background:rgba(244,63,94,0.1);border:1px solid rgba(244,63,94,0.3);
                    border-radius:100px;font-size:0.7rem;font-weight:600;color:#FDA4AF;margin-top:0.5rem;">
          ⚠ Keys Required
        </div>
        """, unsafe_allow_html=True)

    st.markdown("<hr>", unsafe_allow_html=True)

    # ── WORKSPACE NAVIGATION ───────────────────
    st.markdown('<p style="font-size:0.68rem;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;color:#4E5E7A;padding:0 0 0.5rem;">WORKSPACE</p>', unsafe_allow_html=True)
    st.markdown('''<style>
    div[role="radiogroup"] > label > div:first-child { display: none; }
    div[role="radiogroup"] > label { padding: 0.5rem 1rem; border-radius: 8px; margin-bottom: 0.2rem; cursor: pointer; transition: 0.2s; }
    div[role="radiogroup"] > label:hover { background: rgba(99,102,241,0.1); }
    div[role="radiogroup"] > label[data-checked="true"] { background: rgba(99,102,241,0.2); font-weight: bold; border-left: 3px solid #6366F1; color: white !important;}
    </style>''', unsafe_allow_html=True)

    # Apply pending navigation target BEFORE the radio widget renders
    if "_nav_target" in st.session_state:
        st.session_state.current_page = st.session_state._nav_target
        del st.session_state._nav_target

    page = st.radio("Navigation", ["Command Center", "Meetings", "Intelligence", "Action Hub", "People", "Validation"], key="current_page", label_visibility="collapsed")

    # ── MEETING SELECTOR IN SIDEBAR ────────────
    past_meetings = get_user_meetings(st.session_state.get("user_id"))
    if past_meetings:
        st.markdown("<hr>", unsafe_allow_html=True)
        st.markdown('<p style="font-size:0.68rem;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;color:#4E5E7A;padding:0 0 0.25rem;">SELECT MEETING</p>', unsafe_allow_html=True)
        
        m_labels = [f"📄 {m[1]} ({m[4].strftime('%m/%d') if m[4] else 'Recent'})" for m in past_meetings]
        curr_idx = 0
        if st.session_state.get("active_meeting_id"):
            for idx, m in enumerate(past_meetings):
                if m[0] == st.session_state.active_meeting_id:
                    curr_idx = idx
                    break
        
        sel_label = st.selectbox("Choose meeting", m_labels, index=curr_idx, key="sb_meeting_select", label_visibility="collapsed")
        sel_m = past_meetings[m_labels.index(sel_label)]
        
        if st.session_state.get("active_meeting_id") != sel_m[0]:
            st.session_state.active_meeting_id = sel_m[0]
            st.session_state.transcript = sel_m[2]
            st.session_state.report_summary = sel_m[3]
            m_details = get_meeting_by_id(sel_m[0])
            st.session_state.current_meeting_data = m_details
            st.session_state._nav_target = "Intelligence"
            st.rerun()

        if st.button("➕ Upload New Recording", use_container_width=True):
            st.session_state.transcript = None
            st.session_state.report_summary = None
            st.session_state.active_meeting_id = None
            st.session_state.current_meeting_data = None
            st.session_state._nav_target = "Meetings"
            st.rerun()

    st.markdown("<hr>", unsafe_allow_html=True)

    st.markdown("""
    <p style="font-size:0.68rem;font-weight:700;letter-spacing:0.1em;text-transform:uppercase;
              color:#4E5E7A;margin-bottom:0.75rem;">Pipeline</p>
    <div style="display:flex;flex-direction:column;gap:0.5rem;">
      <div style="display:flex;align-items:center;gap:0.5rem;font-size:0.78rem;color:#94A3B8;">
        <span style="color:#6366F1">◆</span> Whisper — Transcription
      </div>
      <div style="display:flex;align-items:center;gap:0.5rem;font-size:0.78rem;color:#94A3B8;">
        <span style="color:#22D3EE">◆</span> Pyannote — Diarization
      </div>
      <div style="display:flex;align-items:center;gap:0.5rem;font-size:0.78rem;color:#94A3B8;">
        <span style="color:#10B981">◆</span> ChromaDB — RAG Index
      </div>
      <div style="display:flex;align-items:center;gap:0.5rem;font-size:0.78rem;color:#94A3B8;">
        <span style="color:#F59E0B">◆</span> Gemini — Intelligence
      </div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown("<hr>", unsafe_allow_html=True)

    st.markdown("""
    <p style="font-size:0.65rem;color:#4E5E7A;line-height:1.6;text-align:center;">
      AI Meeting Synthesizer<br>
      <span style="color:#6366F1">v2.0</span> · Built with Streamlit
    </p>
    """, unsafe_allow_html=True)
    
    if st.session_state.user_id:
        st.markdown("<hr>", unsafe_allow_html=True)
        st.markdown(f'<p style="font-size:0.75rem; color:var(--text-primary); text-align:center;">Logged in as: <strong>{st.session_state.username}</strong></p>', unsafe_allow_html=True)
        if st.button("Logout", use_container_width=True):
            for key in ["user_id", "username", "transcript", "chroma_collection", "report_summary", "structured_data", "active_meeting_id", "current_meeting_data"]:
                st.session_state[key] = None
            st.session_state["messages"] = []
            st.session_state._nav_target = "Command Center"
            st.rerun()


# ─────────────────────────────────────────────
#  GUARD
# ─────────────────────────────────────────────
if not keys_ok:
    st.markdown("""
    <div style="text-align:center;padding:4rem 2rem;">
      <div style="font-size:3rem;margin-bottom:1rem;">🔑</div>
      <h2 style="color:#F0F4FF;font-weight:700;margin-bottom:0.5rem;">API Keys Required</h2>
      <p style="color:#94A3B8;">Please enter your Gemini API Key and Hugging Face Token in the sidebar to continue.</p>
    </div>
    """, unsafe_allow_html=True)
    st.stop()


# ─────────────────────────────────────────────
#  MODEL CACHING
# ─────────────────────────────────────────────
@st.cache_resource
def load_whisper_model():
    return whisper.load_model("base")

@st.cache_resource
def load_diarization_pipeline(token):
    return Pipeline.from_pretrained("pyannote/speaker-diarization-3.1", token=token)

@st.cache_resource
def load_embedding_model():
    return SentenceTransformer("all-MiniLM-L6-v2")

@st.cache_resource
def get_chroma_client():
    return chromadb.Client()


# ─────────────────────────────────────────────
#  HELPERS
# ─────────────────────────────────────────────
def format_timestamp(seconds):
    mins = int(seconds // 60)
    secs = int(seconds % 60)
    return f"[{mins:02d}:{secs:02d}]"

def merge_transcription_and_diarization(whisper_segments, diarization_result):
    merged_transcript = []
    for seg in whisper_segments:
        start, end, text = seg['start'], seg['end'], seg['text'].strip()
        speaker, max_overlap = "Unknown", 0
        for turn, _, spk in diarization_result.itertracks(yield_label=True):
            overlap = max(0, min(end, turn.end) - max(start, turn.start))
            if overlap > max_overlap:
                max_overlap, speaker = overlap, spk
        merged_transcript.append(f"{format_timestamp(start)} {speaker}: {text}")
    return "\n".join(merged_transcript)

def chunk_text(text, chunk_size=5):
    lines = text.split('\n')
    return ["\n".join(lines[i:i+chunk_size]) for i in range(0, len(lines), chunk_size)]

def ensure_chroma_collection():
    """Ensure a valid ChromaDB collection exists for the current transcript."""
    if st.session_state.get("chroma_collection") is not None:
        return st.session_state.chroma_collection

    transcript = st.session_state.get("transcript")
    if not transcript or not transcript.strip():
        return None

    try:
        chroma_client = get_chroma_client()
        active_id = st.session_state.get("active_meeting_id") or "current"
        coll_name = f"meeting_chunks_{active_id}"

        try:
            coll = chroma_client.get_collection(coll_name)
            if coll and coll.count() > 0:
                st.session_state.chroma_collection = coll
                return coll
        except Exception:
            pass

        try:
            chroma_client.delete_collection(coll_name)
        except Exception:
            pass

        collection = chroma_client.create_collection(coll_name)
        embedder = load_embedding_model()
        chunks = chunk_text(transcript)
        if chunks:
            embeddings = embedder.encode(chunks).tolist()
            collection.add(
                embeddings=embeddings,
                documents=chunks,
                ids=[f"chunk_{i}" for i in range(len(chunks))],
            )
        st.session_state.chroma_collection = collection
        return collection
    except Exception:
        return None

def generate_pdf_report(summary, transcript):
    html_content = f"""
    <html><head>
    <style>
      body {{ font-family: Arial, sans-serif; line-height: 1.6; color: #1a1a2e; }}
      h1, h2, h3 {{ color: #4F46E5; }}
      .summary-block {{ background: #f0f0ff; border-left: 4px solid #6366F1;
                        padding: 1rem; border-radius: 4px; margin: 1rem 0; }}
      .transcript {{ font-family: monospace; white-space: pre-wrap; font-size: 11px;
                     background: #f8f8f8; padding: 1rem; border-radius: 4px; }}
      .meta {{ font-size: 11px; color: #666; border-top: 1px solid #eee; padding-top: 0.5rem; }}
    </style>
    </head><body>
      <h1>🧠 Meeting Intelligence Report</h1>
      <p class="meta">Generated by SynthAI Meeting Synthesizer</p>
      <div class="summary-block">{markdown.markdown(summary)}</div>
      <h2>Full Transcript</h2>
      <div class="transcript">{transcript}</div>
    </body></html>
    """
    pdf_buffer = io.BytesIO()
    pisa.CreatePDF(io.StringIO(html_content), dest=pdf_buffer)
    pdf_buffer.seek(0)
    return pdf_buffer.getvalue()

def generate_docx_report(summary, transcript):
    doc = docx.Document()
    doc.add_heading('Meeting Intelligence Report', 0)
    doc.add_heading('Summary', level=1)
    for line in summary.split('\n'):
        if not line.strip():
            continue
        if line.startswith('# '):
            doc.add_heading(line.replace('# ', ''), level=1)
        elif line.startswith('## '):
            doc.add_heading(line.replace('## ', ''), level=2)
        elif line.startswith('### '):
            doc.add_heading(line.replace('### ', ''), level=3)
        elif line.startswith(('- ', '* ')):
            doc.add_paragraph(line[2:], style='List Bullet')
        else:
            doc.add_paragraph(line)
    doc.add_heading('Full Transcript', level=1)
    for line in transcript.split('\n'):
        if line.strip():
            doc.add_paragraph(line.strip())
    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)
    return docx_buffer.getvalue()

# ─────────────────────────────────────────────
#  AUTHENTICATION
# ─────────────────────────────────────────────
# (Authentication gate is now handled before the sidebar above)
# ─────────────────────────────────────────────
#  ROUTER LOGIC
# ─────────────────────────────────────────────

# ═══════ COMMAND CENTER ═══════
if page == "Command Center":
    st.markdown("""
    <div class="hero-container">
      <div class="hero-logo">
        <div class="hero-logo-icon">🧠</div>
        <div class="hero-brand">SynthAI</div>
      </div>
      <div class="hero-tagline">Meeting Intelligence Platform — Transcribe, Analyze, Synthesize</div>
      <div class="hero-badges">
        <span class="hero-badge badge-whisper">⚡ Whisper Transcription</span>
        <span class="hero-badge badge-gemini">✦ Gemini Intelligence</span>
        <span class="hero-badge badge-rag">◈ RAG Chat</span>
      </div>
    </div>
    """, unsafe_allow_html=True)

    # Quick status cards
    has_meeting = st.session_state.transcript is not None
    if has_meeting:
        word_count = len(st.session_state.transcript.split())
        lines_count = len(st.session_state.transcript.strip().split('\n'))
        st.markdown("<hr>", unsafe_allow_html=True)
        st.markdown("### 📊 Active Meeting Session")
        c1, c2, c3 = st.columns(3)
        c1.metric("Status", "✅ Meeting Loaded")
        c2.metric("Utterances", f"{lines_count}")
        c3.metric("Words", f"{word_count:,}")
        
        col_btn1, col_btn2 = st.columns(2)
        with col_btn1:
            if st.button("🔍 View Meeting Intelligence & Report", use_container_width=True, type="primary"):
                st.session_state._nav_target = "Intelligence"
                st.rerun()
        with col_btn2:
            if st.button("➕ Upload Another Meeting", use_container_width=True):
                st.session_state._nav_target = "Meetings"
                st.rerun()
    else:
        st.markdown("<hr>", unsafe_allow_html=True)
        st.info("👋 Welcome! Select an existing meeting from the sidebar or click below to upload a new recording.")
        if st.button("🎙️ Upload & Process New Meeting", use_container_width=True, type="primary"):
            st.session_state._nav_target = "Meetings"
            st.rerun()

    # Meeting History from database
    if st.session_state.get("user_id"):
        past_meetings = get_user_meetings(st.session_state.user_id)
        if past_meetings:
            st.markdown("<hr>", unsafe_allow_html=True)
            st.markdown("### 📂 Your Saved Meetings")
            for m_id, m_file, m_transcript, m_summary, m_date in past_meetings:
                date_str = m_date.strftime("%Y-%m-%d %H:%M") if m_date else "Recent"
                word_ct = len(m_transcript.split()) if m_transcript else 0
                is_current = (st.session_state.get("active_meeting_id") == m_id)
                tag = " (Currently Active)" if is_current else ""
                with st.expander(f"📄 {m_file} — {date_str} ({word_ct:,} words){tag}"):
                    if m_summary:
                        st.markdown("**Executive Summary:**")
                        st.markdown(m_summary[:400] + ("..." if len(m_summary) > 400 else ""))
                    if not is_current:
                        if st.button(f"Load this meeting", key=f"cmd_load_{m_id}"):
                            st.session_state.active_meeting_id = m_id
                            st.session_state.transcript = m_transcript
                            st.session_state.report_summary = m_summary
                            m_details = get_meeting_by_id(m_id)
                            st.session_state.current_meeting_data = m_details
                            st.session_state._nav_target = "Intelligence"
                            st.rerun()
                    else:
                        st.success("Currently active. Navigate to Intelligence or Action Hub to view details.")

    st.stop()

# ═══════ ACTION HUB ═══════
elif page == "Action Hub":
    st.markdown("""
    <div class="section-header">
      <div class="section-header-icon icon-amber">⚡</div>
      <div class="section-header-text">
        <h3>Action Hub</h3>
        <span>Action items and deliverables extracted from your meeting</span>
      </div>
    </div>
    """, unsafe_allow_html=True)

    if not st.session_state.transcript:
        st.warning("No meeting loaded yet. Please select a meeting from the sidebar or go to **Meetings** to upload a recording.")
        if st.button("Go to Meetings"):
            st.session_state._nav_target = "Meetings"
            st.rerun()
        st.stop()

    structured_actions = []
    if st.session_state.get("current_meeting_data") and st.session_state.current_meeting_data.get("action_items"):
        structured_actions = st.session_state.current_meeting_data["action_items"]
    elif st.session_state.get("active_meeting_id"):
        m_details = get_meeting_by_id(st.session_state.active_meeting_id)
        if m_details and m_details.get("action_items"):
            structured_actions = m_details["action_items"]
            st.session_state.current_meeting_data = m_details

    if structured_actions:
        st.markdown(f"**{len(structured_actions)}** action items identified:")
        for idx, item in enumerate(structured_actions):
            status_key = f"action_{st.session_state.get('active_meeting_id', 0)}_{idx}"
            if status_key not in st.session_state:
                st.session_state[status_key] = (item.get("status") == "Completed")

            col_check, col_text = st.columns([0.05, 0.95])
            with col_check:
                st.session_state[status_key] = st.checkbox("", value=st.session_state[status_key], key=f"cb_{status_key}", label_visibility="collapsed")
            with col_text:
                priority = item.get("priority") or "Medium"
                priority_color = "#F43F5E" if priority.lower() == "high" else ("#F59E0B" if priority.lower() == "medium" else "#10B981")
                assignee = item.get("assigned_participant") or "Unassigned"
                deadline = f" · ⏰ {item.get('deadline')}" if item.get("deadline") else ""
                
                desc_text = item.get('description', '')
                if st.session_state[status_key]:
                    st.markdown(f'<p style="color:#94A3B8;text-decoration:line-through;margin:0;">{desc_text}</p>', unsafe_allow_html=True)
                else:
                    st.markdown(f'<p style="color:#F0F4FF;font-weight:500;margin:0;">{desc_text}</p>', unsafe_allow_html=True)
                st.markdown(f'<span style="font-size:0.7rem;color:#818CF8;">👤 {assignee}</span> <span style="font-size:0.7rem;color:{priority_color};font-weight:600;margin-left:0.5rem;">● {priority}</span><span style="font-size:0.7rem;color:#94A3B8;">{deadline}</span>', unsafe_allow_html=True)

        st.markdown("<hr>", unsafe_allow_html=True)
        done_count = sum(1 for i in range(len(structured_actions)) if st.session_state.get(f"action_{st.session_state.get('active_meeting_id', 0)}_{i}", False))
        total = len(structured_actions)
        pct = done_count / total if total > 0 else 0
        st.progress(pct)
        st.markdown(f"**{done_count}/{total}** action items completed")
    else:
        report = st.session_state.get("report_summary", "")
        action_lines = []
        in_action_section = False
        for line in report.split('\n'):
            lower = line.lower().strip()
            if 'action item' in lower or 'action' in lower and 'owner' in lower:
                in_action_section = True
                continue
            if in_action_section:
                if line.strip().startswith(('#', '##', '###')) and 'action' not in line.lower():
                    in_action_section = False
                    continue
                stripped = line.strip().lstrip('-').lstrip('*').lstrip('•').strip()
                if stripped and len(stripped) > 3:
                    action_lines.append(stripped)

        if not action_lines:
            for line in report.split('\n'):
                stripped = line.strip().lstrip('-').lstrip('*').lstrip('•').strip()
                for keyword in ['should', 'will', 'needs to', 'must', 'assigned to', 'responsible', 'follow up', 'deadline']:
                    if keyword in stripped.lower() and len(stripped) > 10:
                        action_lines.append(stripped)
                        break

        if action_lines:
            for idx, item in enumerate(action_lines):
                status_key = f"action_{idx}"
                if status_key not in st.session_state:
                    st.session_state[status_key] = False

                col_check, col_text = st.columns([0.05, 0.95])
                with col_check:
                    st.session_state[status_key] = st.checkbox("", value=st.session_state[status_key], key=f"cb_{status_key}", label_visibility="collapsed")
                with col_text:
                    if st.session_state[status_key]:
                        st.markdown(f'<p style="color:#94A3B8;text-decoration:line-through;">{item}</p>', unsafe_allow_html=True)
                    else:
                        st.markdown(f'<p style="color:#F0F4FF;">{item}</p>', unsafe_allow_html=True)

            st.markdown("<hr>", unsafe_allow_html=True)
            done_count = sum(1 for i in range(len(action_lines)) if st.session_state.get(f"action_{i}", False))
            total = len(action_lines)
            pct = done_count / total if total > 0 else 0
            st.progress(pct)
            st.markdown(f"**{done_count}/{total}** action items completed")
        else:
            st.info("No specific action items found in the current meeting.")

    st.stop()

# ═══════ PEOPLE ═══════
elif page == "People":
    st.markdown("""
    <div class="section-header">
      <div class="section-header-icon icon-emerald">👥</div>
      <div class="section-header-text">
        <h3>People & Participants</h3>
        <span>Speaker participation analytics from diarized transcript</span>
      </div>
    </div>
    """, unsafe_allow_html=True)

    if not st.session_state.transcript:
        st.warning("No meeting loaded yet. Please select a meeting from the sidebar or go to **Meetings** to upload a recording.")
        if st.button("Go to Meetings"):
            st.session_state._nav_target = "Meetings"
            st.rerun()
        st.stop()

    speaker_data = {}
    transcript_lines = st.session_state.transcript.strip().split('\n')
    for line in transcript_lines:
        parts = line.split(' ', 2)
        if len(parts) >= 3:
            speaker = parts[1].rstrip(':')
            text = parts[2] if len(parts) > 2 else ""
            word_count = len(text.split())
            if speaker not in speaker_data:
                speaker_data[speaker] = {"utterances": 0, "words": 0}
            speaker_data[speaker]["utterances"] += 1
            speaker_data[speaker]["words"] += word_count

    if not speaker_data:
        st.info("No speaker data found in the transcript.")
        st.stop()

    total_words = sum(v["words"] for v in speaker_data.values())
    total_utterances = sum(v["utterances"] for v in speaker_data.values())

    c1, c2, c3 = st.columns(3)
    with c1:
        st.markdown(f"""
        <div class="metric-card">
          <div class="metric-icon icon-emerald">👥</div>
          <div>
            <div class="metric-label">Speakers</div>
            <div class="metric-value">{len(speaker_data)}</div>
          </div>
        </div>""", unsafe_allow_html=True)
    with c2:
        st.markdown(f"""
        <div class="metric-card">
          <div class="metric-icon icon-indigo">🎙</div>
          <div>
            <div class="metric-label">Total Utterances</div>
            <div class="metric-value">{total_utterances}</div>
          </div>
        </div>""", unsafe_allow_html=True)
    with c3:
        st.markdown(f"""
        <div class="metric-card">
          <div class="metric-icon icon-cyan">💬</div>
          <div>
            <div class="metric-label">Total Words</div>
            <div class="metric-value">{total_words:,}</div>
          </div>
        </div>""", unsafe_allow_html=True)

    st.markdown('<div style="height:1.5rem;"></div>', unsafe_allow_html=True)

    colors = ["#6366F1", "#22D3EE", "#10B981", "#F59E0B", "#F43F5E", "#A78BFA", "#FB923C"]
    for idx, (speaker, data) in enumerate(sorted(speaker_data.items())):
        color = colors[idx % len(colors)]
        pct = (data["words"] / total_words * 100) if total_words > 0 else 0

        st.markdown(f"""
        <div class="glass-card" style="margin-bottom:1rem; padding:1.25rem 1.5rem;">
          <div style="display:flex;justify-content:space-between;align-items:center;margin-bottom:0.75rem;">
            <div style="display:flex;align-items:center;gap:0.75rem;">
              <div style="width:40px;height:40px;border-radius:50%;background:{color};display:flex;align-items:center;justify-content:center;font-size:1.1rem;font-weight:700;color:white;">
                {speaker[0] if speaker else '?'}
              </div>
              <div>
                <div style="font-weight:700;color:#F0F4FF;font-size:1rem;">{speaker}</div>
                <div style="color:#94A3B8;font-size:0.75rem;">{data['utterances']} utterances · {data['words']} words</div>
              </div>
            </div>
            <div style="font-size:1.5rem;font-weight:800;color:{color};">{pct:.1f}%</div>
          </div>
          <div style="width:100%;height:8px;background:rgba(255,255,255,0.05);border-radius:4px;overflow:hidden;">
            <div style="width:{pct}%;height:100%;background:{color};border-radius:4px;transition:width 0.5s;"></div>
          </div>
        </div>
        """, unsafe_allow_html=True)

    st.stop()

# ═══════ VALIDATION ═══════
elif page == "Validation":
    st.markdown("""
    <div class="section-header">
      <div class="section-header-icon icon-emerald">✅</div>
      <div class="section-header-text">
        <h3>Validation Statistics</h3>
        <span>Compare Whisper transcription against a reference text</span>
      </div>
    </div>
    """, unsafe_allow_html=True)

    if not st.session_state.transcript:
        st.warning("No meeting loaded yet. Please select a meeting from the sidebar or go to **Meetings** to upload a recording.")
        if st.button("Go to Meetings"):
            st.session_state._nav_target = "Meetings"
            st.rerun()
        st.stop()

    whisper_words = []
    for line in st.session_state.transcript.strip().split('\n'):
        cleaned = re.sub(r'\[\d+:\d+\]', '', line)
        cleaned = re.sub(r'SPEAKER_\d+:', '', cleaned, flags=re.IGNORECASE)
        cleaned = re.sub(r'Speaker\s*\d*:', '', cleaned, flags=re.IGNORECASE)
        words = cleaned.lower().split()
        whisper_words.extend(words)

    st.markdown("#### Upload Reference Text")
    st.markdown('<p style="color:#94A3B8;font-size:0.85rem;">Upload a <code>.txt</code> file containing the ground-truth transcript to compare against the Whisper output.</p>', unsafe_allow_html=True)

    ref_file = st.file_uploader("Reference Transcript (.txt)", type=["txt"], label_visibility="collapsed")

    if ref_file is not None:
        ref_text = ref_file.read().decode("utf-8", errors="ignore")
        cleaned_ref = re.sub(r'\[\d+:\d+\]', '', ref_text)
        cleaned_ref = re.sub(r'SPEAKER_\d+:', '', cleaned_ref, flags=re.IGNORECASE)
        cleaned_ref = re.sub(r'Speaker\s*\d*:', '', cleaned_ref, flags=re.IGNORECASE)
        ref_words = cleaned_ref.lower().split()

        ref_set = {}
        for w in ref_words:
            ref_set[w] = ref_set.get(w, 0) + 1

        whisper_set = {}
        for w in whisper_words:
            whisper_set[w] = whisper_set.get(w, 0) + 1

        correct = 0
        for w, count in ref_set.items():
            correct += min(count, whisper_set.get(w, 0))

        ref_total = len(ref_words)
        whisper_total = len(whisper_words)
        accuracy = (correct / ref_total * 100) if ref_total > 0 else 0
        target = 90.0
        passed = accuracy >= target

        missing_words = {}
        for w, count in ref_set.items():
            diff = count - whisper_set.get(w, 0)
            if diff > 0:
                missing_words[w] = diff

        extra_words = {}
        for w, count in whisper_set.items():
            diff = count - ref_set.get(w, 0)
            if diff > 0:
                extra_words[w] = diff

        st.markdown("<hr>", unsafe_allow_html=True)

        col1, col2 = st.columns([1, 2])
        with col1:
            color = "#10B981" if passed else "#F43F5E"
            st.markdown(f'''
            <div class="glass-card" style="text-align:center; padding: 2rem;">
              <h4 style="color:#94A3B8;font-size:0.8rem;text-transform:uppercase;margin-bottom:0.5rem;">Transcription Accuracy</h4>
              <h1 style="font-size:3.5rem;margin:0;color:{color};">{accuracy:.2f}%</h1>
              <p style="color:#94A3B8;font-size:0.8rem;margin-top:0.5rem;">Target: {target:.0f}%</p>
            </div>''', unsafe_allow_html=True)
        with col2:
            if passed:
                st.success(f"PASS — transcription meets the {target:.0f}% target.")
            else:
                st.error(f"FAIL — transcription is below the {target:.0f}% target.")

        st.markdown("<hr>", unsafe_allow_html=True)
        st.markdown("## Validation Statistics")

        c1, c2, c3 = st.columns(3)
        c1.metric("Reference Words", f"{ref_total:,}")
        c2.metric("Whisper Words", f"{whisper_total:,}")
        c3.metric("Correct Words", f"{correct:,}")

        st.markdown("<br>", unsafe_allow_html=True)

        if missing_words:
            top_missing = sorted(missing_words.items(), key=lambda x: -x[1])[:20]
            missing_str = ", ".join([f"**{w}** (×{c})" for w, c in top_missing])
            st.warning(f"**{sum(missing_words.values())}** missing words detected. Top: {missing_str}")
        else:
            st.success("No missing words detected.")

        if extra_words:
            top_extra = sorted(extra_words.items(), key=lambda x: -x[1])[:20]
            extra_str = ", ".join([f"**{w}** (×{c})" for w, c in top_extra])
            st.warning(f"**{sum(extra_words.values())}** extra words detected. Top: {extra_str}")
        else:
            st.success("No extra words detected.")

    else:
        st.info("Upload a reference text file above to begin validation.")

    st.stop()

# ═══════ INTELLIGENCE GUARD ═══════
elif page == "Intelligence":
    if not st.session_state.transcript:
        st.warning("No meeting loaded yet. Please select an existing meeting from the sidebar or go to **Meetings** to upload a recording.")
        if st.button("🎙️ Go to Meetings to Upload"):
            st.session_state._nav_target = "Meetings"
            st.rerun()
        st.stop()

# ═══════ MEETINGS (UPLOAD) ═══════
if page == "Meetings":

    st.markdown("""
    <div style="max-width:680px;margin:0 auto 0.5rem;">
      <div class="upload-zone-wrapper">
        <span class="upload-icon">🎙️</span>
        <div class="upload-title">Drop your meeting recording here</div>
        <div class="upload-subtitle">Upload any audio or video file to begin AI-powered analysis</div>
        <div class="format-chips">
          <span class="format-chip">.mp3</span>
          <span class="format-chip">.wav</span>
          <span class="format-chip">.m4a</span>
          <span class="format-chip">.mp4</span>
        </div>
      </div>
    </div>
    """, unsafe_allow_html=True)

    col_center = st.columns([1, 2, 1])[1]
    with col_center:
        uploaded_file = st.file_uploader(
            "Upload meeting recording",
            type=["mp3", "wav", "m4a", "mp4"],
            label_visibility="collapsed",
        )

    if uploaded_file is not None:
        # File info card
        file_size_mb = uploaded_file.size / (1024 * 1024)
        st.markdown(f"""
        <div style="max-width:680px;margin:0.75rem auto 1.25rem;
                    background:rgba(99,102,241,0.08);border:1px solid rgba(99,102,241,0.2);
                    border-radius:12px;padding:1rem 1.25rem;
                    display:flex;align-items:center;gap:0.75rem;">
          <span style="font-size:1.6rem;">🎵</span>
          <div style="flex:1;">
            <div style="font-weight:600;font-size:0.88rem;color:#F0F4FF;">{uploaded_file.name}</div>
            <div style="font-size:0.75rem;color:#94A3B8;margin-top:0.15rem;">{file_size_mb:.2f} MB · Ready for processing</div>
          </div>
          <div style="background:rgba(16,185,129,0.15);border:1px solid rgba(16,185,129,0.3);
                      border-radius:100px;padding:0.25rem 0.65rem;font-size:0.68rem;
                      font-weight:700;color:#6EE7B7;text-transform:uppercase;">Loaded</div>
        </div>
        """, unsafe_allow_html=True)

        col_btn = st.columns([1, 2, 1])[1]
        with col_btn:
            process_btn = st.button("🚀  Analyse Meeting", use_container_width=True)

        if process_btn:
            file_ext = os.path.splitext(uploaded_file.name)[1].lower()

            # ── STEP TRACKER ──────────────────────
            step_placeholder = st.empty()

            def render_steps(active_idx):
                steps = [
                    ("🎙", "Transcribing"),
                    ("👥", "Diarizing"),
                    ("🗃", "Indexing"),
                    ("✦", "Synthesizing"),
                ]
                items_html = ""
                for i, (icon, label) in enumerate(steps):
                    if i < active_idx:
                        dot_cls, lbl_cls = "done", "done"
                        icon_inner = "✓"
                    elif i == active_idx:
                        dot_cls, lbl_cls = "active", "active"
                        icon_inner = icon
                    else:
                        dot_cls, lbl_cls = "waiting", ""
                        icon_inner = icon
                    items_html += f"""
                    <div class="step-item {'done' if i < active_idx else ('active' if i == active_idx else '')}">
                      <div class="step-dot {dot_cls}">{icon_inner}</div>
                      <div class="step-label {lbl_cls}">{label}</div>
                    </div>
                    """
                step_placeholder.markdown(
                    f'<div class="steps-container">{items_html}</div>',
                    unsafe_allow_html=True,
                )

            try:
                with tempfile.NamedTemporaryFile(delete=False, suffix=file_ext) as tmp:
                    tmp.write(uploaded_file.read())
                    tmp_path = tmp.name

                render_steps(0)
                whisper_model = load_whisper_model()
                diarization_pipeline = load_diarization_pipeline(hf_token)

                try:
                    audio_array = whisper.load_audio(tmp_path, sr=16000)
                    if len(audio_array) == 0:
                        st.error("The uploaded file does not contain audio (or the audio track is empty).")
                        st.stop()
                except Exception as e:
                    st.error(f"Could not extract audio: {e}")
                    st.stop()

                with st.spinner("Transcribing speech with Whisper…"):
                    result = whisper_model.transcribe(audio_array)

                render_steps(1)

                with st.spinner("Identifying speakers with Pyannote…"):
                    waveform_tensor = torch.from_numpy(audio_array).unsqueeze(0)
                    diarization = diarization_pipeline({"waveform": waveform_tensor, "sample_rate": 16000})

                render_steps(2)

                diarization_annotation = getattr(diarization, "speaker_diarization", diarization)
                final_transcript = merge_transcription_and_diarization(result['segments'], diarization_annotation)

                if not final_transcript:
                    st.error("Transcription returned empty output. Please try a different file.")
                    st.stop()

                st.session_state.transcript = final_transcript

                with st.spinner("Building RAG index in ChromaDB…"):
                    chroma_client = get_chroma_client()
                    try:
                        chroma_client.delete_collection("meeting_chunks")
                    except Exception:
                        pass
                    collection = chroma_client.create_collection("meeting_chunks")
                    embedder = load_embedding_model()
                    chunks = chunk_text(final_transcript)
                    embeddings = embedder.encode(chunks).tolist()
                    collection.add(
                        embeddings=embeddings,
                        documents=chunks,
                        ids=[f"chunk_{i}" for i in range(len(chunks))],
                    )
                    st.session_state.chroma_collection = collection

                render_steps(3)

                with st.spinner("Generating structured intelligence with Gemini…"):
                    client = genai.Client(api_key=gemini_key)
                    from llm_service import process_transcript_with_llm
                    from database import get_db, SessionLocal
                    from models import Meeting, Participant, ActionItem, KeyPoint, KeyDecision
                    
                    try:
                        structured_data = process_transcript_with_llm(client, final_transcript)
                        st.session_state.structured_data = structured_data
                        
                        # Save to database
                        db = SessionLocal()
                        
                        # Find or create participants
                        participant_objs = []
                        for p_name in structured_data.participants:
                            p = db.query(Participant).filter(Participant.name == p_name).first()
                            if not p:
                                p = Participant(name=p_name)
                                db.add(p)
                            participant_objs.append(p)
                        db.commit()
                        
                        # Create Meeting
                        meeting = Meeting(
                            filename=uploaded_file.name,
                            transcript=final_transcript,
                            summary=structured_data.summary,
                            user_id=st.session_state.user_id
                        )
                        meeting.participants.extend(participant_objs)
                        db.add(meeting)
                        db.commit()
                        db.refresh(meeting)
                        
                        # Add Action Items
                        for ai in structured_data.action_items:
                            p_id = None
                            if ai.assigned_participant:
                                p = db.query(Participant).filter(Participant.name == ai.assigned_participant).first()
                                if p:
                                    p_id = p.id
                            action = ActionItem(
                                meeting_id=meeting.id,
                                participant_id=p_id,
                                description=ai.description,
                                deadline=ai.deadline,
                                priority=ai.priority,
                                status=ai.status
                            )
                            db.add(action)
                            
                        # Add Key Points
                        for kp in structured_data.key_points:
                            db.add(KeyPoint(meeting_id=meeting.id, point=kp))
                            
                        # Add Key Decisions
                        for kd in structured_data.decisions:
                            db.add(KeyDecision(meeting_id=meeting.id, decision=kd))
                            
                        db.commit()
                        # Extract meeting ID before closing the session
                        saved_meeting_id = meeting.id
                        db.close()
                        
                        # Fallback simple string summary
                        st.session_state.report_summary = structured_data.summary
                        
                    except Exception as e:
                        saved_meeting_id = None
                        st.session_state.report_summary = f"⚠ Failed to generate structured intelligence: {e}"
                        st.session_state.structured_data = None

                render_steps(4)  # all done
                st.session_state.active_meeting_id = saved_meeting_id
                st.session_state._nav_target = "Intelligence"
                st.success("✅  Analysis complete! Navigating to Intelligence...")
                import time; time.sleep(0.8)
                st.rerun()

            except Exception as e:
                st.error(f"An unexpected error occurred: {e}")
            finally:
                if 'tmp_path' in locals() and os.path.exists(tmp_path):
                    try:
                        os.remove(tmp_path)
                    except Exception:
                        pass
        st.stop()
    st.stop()


# ─────────────────────────────────────────────
#  RESULTS VIEW (INTELLIGENCE)
# ─────────────────────────────────────────────
if page == "Intelligence" and st.session_state.transcript:

    transcript_lines = st.session_state.transcript.strip().split('\n')
    speakers = set()
    for line in transcript_lines:
        parts = line.split(' ')
        if len(parts) > 1:
            spk = [p.rstrip(':') for p in parts if p.startswith('SPEAKER')]
            speakers.update(spk)

    # ── STATS ROW ──────────────────────────────
    st.markdown('<div style="height:0.5rem;"></div>', unsafe_allow_html=True)
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        st.markdown(f"""
        <div class="metric-card">
          <div class="metric-icon icon-indigo">🎙</div>
          <div>
            <div class="metric-label">Utterances</div>
            <div class="metric-value">{len(transcript_lines)}</div>
          </div>
        </div>""", unsafe_allow_html=True)
    with c2:
        word_count = len(st.session_state.transcript.split())
        st.markdown(f"""
        <div class="metric-card">
          <div class="metric-icon icon-cyan">💬</div>
          <div>
            <div class="metric-label">Words</div>
            <div class="metric-value">{word_count:,}</div>
          </div>
        </div>""", unsafe_allow_html=True)
    with c3:
        st.markdown(f"""
        <div class="metric-card">
          <div class="metric-icon icon-emerald">👥</div>
          <div>
            <div class="metric-label">Speakers</div>
            <div class="metric-value">{max(len(speakers), 1)}</div>
          </div>
        </div>""", unsafe_allow_html=True)
    with c4:
        report_status = "Ready" if st.session_state.report_summary else "Pending"
        st.markdown(f"""
        <div class="metric-card">
          <div class="metric-icon icon-amber">✦</div>
          <div>
            <div class="metric-label">Report</div>
            <div class="metric-value">{report_status}</div>
          </div>
        </div>""", unsafe_allow_html=True)

    st.markdown('<div style="height:1rem;"></div>', unsafe_allow_html=True)

    # ── TABS ───────────────────────────────────
    tab1, tab2 = st.tabs(["📋  Meeting Intelligence", "💬  AI Chat Assistant"])

    # ════════════════════════════════════════════
    #  TAB 1 — REPORT
    # ════════════════════════════════════════════
    with tab1:

        if st.session_state.report_summary:

            # Executive Summary card
            st.markdown("""
            <div class="section-header">
              <div class="section-header-icon icon-indigo">✦</div>
              <div class="section-header-text">
                <h3>Meeting Intelligence Report</h3>
                <span>AI-generated analysis powered by Gemini</span>
              </div>
            </div>
            """, unsafe_allow_html=True)

            st.markdown('<div class="report-executive">', unsafe_allow_html=True)
            if st.session_state.structured_data:
                sd = st.session_state.structured_data
                st.markdown(f"**Executive Summary**\n\n{sd.summary}")
                
                if sd.key_points:
                    st.markdown("**Key Points**")
                    for kp in sd.key_points:
                        st.markdown(f"- {kp}")
                        
                if sd.decisions:
                    st.markdown("**Key Decisions**")
                    for kd in sd.decisions:
                        st.markdown(f"- {kd}")
                        
                if sd.action_items:
                    st.markdown("**Action Items**")
                    for ai in sd.action_items:
                        st.markdown(f"- {ai.description} (Assigned to: {ai.assigned_participant}, Deadline: {ai.deadline}, Priority: {ai.priority}, Status: {ai.status})")
            else:
                st.markdown(st.session_state.report_summary)
            st.markdown('</div>', unsafe_allow_html=True)

            st.markdown('<div style="height:0.5rem;"></div>', unsafe_allow_html=True)

            # Export section
            st.markdown("""
            <div class="export-header">
              ↓ &nbsp;Export Report
            </div>
            """, unsafe_allow_html=True)

            exp_col1, exp_col2, exp_col3, exp_col4 = st.columns(4)

            with exp_col1:
                pdf_bytes = generate_pdf_report(
                    st.session_state.report_summary, st.session_state.transcript
                )
                st.download_button(
                    label="📄  Download PDF",
                    data=pdf_bytes,
                    file_name="Meeting_Report.pdf",
                    mime="application/pdf",
                    use_container_width=True,
                    key="dl_pdf",
                )

            with exp_col2:
                docx_bytes = generate_docx_report(
                    st.session_state.report_summary, st.session_state.transcript
                )
                st.download_button(
                    label="📝  Download DOCX",
                    data=docx_bytes,
                    file_name="Meeting_Report.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True,
                    key="dl_docx",
                )

            with exp_col3:
                summary_txt = (st.session_state.report_summary or "").encode("utf-8")
                st.download_button(
                    label="📋  Download Summary (.txt)",
                    data=summary_txt,
                    file_name="meeting_summary.txt",
                    mime="text/plain",
                    use_container_width=True,
                    key="dl_summary_txt",
                )

            with exp_col4:
                transcript_txt = (st.session_state.transcript or "").encode("utf-8")
                st.download_button(
                    label="📃  Download Transcript (.txt)",
                    data=transcript_txt,
                    file_name="diarized_transcript.txt",
                    mime="text/plain",
                    use_container_width=True,
                    key="dl_transcript",
                )

        st.markdown("<hr>", unsafe_allow_html=True)

        # Full Transcript
        st.markdown("""
        <div class="section-header">
          <div class="section-header-icon icon-cyan">📜</div>
          <div class="section-header-text">
            <h3>Full Diarized Transcript</h3>
            <span>Speaker-attributed, timestamped</span>
          </div>
        </div>
        """, unsafe_allow_html=True)

        st.code(st.session_state.transcript, language="text")

        st.markdown("<hr>", unsafe_allow_html=True)

        # New file button
        st.markdown('<div class="new-file-btn">', unsafe_allow_html=True)
        if st.button("↺  Process a New File", use_container_width=False):
            for key in ["transcript", "chroma_collection", "report_summary", "structured_data", "messages", "active_meeting_id", "current_meeting_data"]:
                st.session_state[key] = None if key != "messages" else []
            st.session_state._nav_target = "Meetings"
            st.rerun()
        st.markdown('</div>', unsafe_allow_html=True)

    # ════════════════════════════════════════════
    #  TAB 2 — CHAT
    # ════════════════════════════════════════════
    with tab2:

        st.markdown("""
        <div class="section-header">
          <div class="section-header-icon icon-indigo">💬</div>
          <div class="section-header-text">
            <h3>AI Meeting Assistant</h3>
            <span>Ask anything about your meeting — powered by RAG + Gemini</span>
          </div>
        </div>
        """, unsafe_allow_html=True)

        # Chat history
        for msg in st.session_state.messages:
            role = msg["role"]
            with st.chat_message(role, avatar="🧑" if role == "user" else "🧠"):
                st.markdown(msg["content"])

        # Chat input
        if prompt := st.chat_input("Ask about the meeting… e.g. 'What were the key action items?'"):
            st.session_state.messages.append({"role": "user", "content": prompt})
            with st.chat_message("user", avatar="🧑"):
                st.markdown(prompt)

            with st.chat_message("assistant", avatar="🧠"):
                msg_placeholder = st.empty()

                context = ""
                collection = ensure_chroma_collection()
                if collection is not None:
                    try:
                        embedder = load_embedding_model()
                        query_embedding = embedder.encode([prompt]).tolist()
                        count = collection.count()
                        if count > 0:
                            results = collection.query(
                                query_embeddings=query_embedding,
                                n_results=min(3, count),
                            )
                            if results and results.get('documents') and results['documents'][0]:
                                context = "\n\n".join(results['documents'][0])
                    except Exception:
                        context = ""

                # Resilient fallback to transcript if RAG collection query was empty
                if not context and st.session_state.get("transcript"):
                    context = st.session_state.transcript[:6000]

                if not context:
                    context = "No meeting transcript available."

                system_prompt = f"""You are an expert AI meeting assistant. Answer the user's question based ONLY on the provided meeting transcript context below.
If the information is not present in the context, say: "I don't have enough information from this meeting to answer that."
Be concise, professional, and precise.

Context:
{context}
"""
                if not gemini_key:
                    st.error("Please provide a Gemini API Key in the sidebar.")
                else:
                    try:
                        client = genai.Client(api_key=gemini_key)
                        import time
                        response = None
                        for attempt in range(3):
                            try:
                                response = client.models.generate_content_stream(
                                    model='gemini-2.5-flash',
                                    contents=system_prompt + f"\nUser Question: {prompt}",
                                )
                                break
                            except Exception as e:
                                if "503" in str(e) and attempt < 2:
                                    time.sleep(2)
                                else:
                                    raise e

                        full_response = ""
                        if response:
                            for chunk in response:
                                if chunk.text:
                                    full_response += chunk.text
                                    msg_placeholder.markdown(full_response + " ▌")

                        msg_placeholder.markdown(full_response)
                        st.session_state.messages.append({"role": "assistant", "content": full_response})

                    except Exception as e:
                        st.error(f"Could not generate response: {e}")
